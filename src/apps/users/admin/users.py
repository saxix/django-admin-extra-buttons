import logging
import io
import openpyxl

from django.contrib import admin, messages
from django.contrib.auth.admin import UserAdmin
from django.utils.safestring import mark_safe
from django.utils.translation import gettext_lazy as _
from django.utils.translation import gettext
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Count, Q, Prefetch
from django.http import HttpResponse, HttpResponseNotAllowed
from django.utils import timezone
from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.shortcuts import render, redirect
from rest_framework.reverse import reverse_lazy

# ✅ ИСПОЛЬЗУЕМ ОБЪЕДИНЕННЫЙ МОДУЛЬ ВМЕСТО ОРИГИНАЛЬНОГО admin_extra_buttons
# from admin_extra_buttons.decorators import button
# from admin_extra_buttons.mixins import ExtraButtonsMixin
import sys
import os
sys.path.append(os.path.join(os.path.dirname(__file__), '../../../..'))
from django_admin_buttons import ExtraButtonsMixin, button

from apps.users.helpers.users_approving import UserApproveHelper, UserHelper
from apps.users.models import BlackListedAccessToken, ContractOffer
from apps.users.models import User
from apps.users.forms import UserCreationForm
from apps.users.choices import UserRole
from apps.agents.services import Sender
from apps.agents_requests.models import AgentsRequestAttachment
from apps.agents.choices import AgentsStatus
from apps.agents.models import Agent

logger = logging.getLogger(__name__)


class AgentsInline(admin.TabularInline):
    verbose_name = _("Agent")
    verbose_name_plural = _("Agents")
    related_name = "agents"
    model = User.agents.through
    fields = (
        "agent",
        "get_inn",
        "get_name",
        "get_status",
        "get_created_at",
        "get_uuid_1c",
        "get_is_deleted",
    )
    readonly_fields = (
        "get_inn",
        "get_name",
        "get_status",
        "get_created_at",
        "get_uuid_1c",
        "get_is_deleted",
    )
    can_delete = True
    show_change_link = True
    extra = 0

    def get_inn(self, instance):
        return instance.agent.inn

    get_inn.short_description = "inn"

    def get_name(self, instance):
        return instance.agent.name

    get_name.short_description = _("Name")

    def get_status(self, instance):
        return instance.agent.status

    get_status.short_description = _("Status")

    def get_created_at(self, instance):
        return instance.agent.created_at

    get_created_at.short_description = _("created at")

    def get_uuid_1c(self, instance):
        return instance.agent.uuid_1c

    get_uuid_1c.short_description = _("UUID 1C")

    def get_is_deleted(self, instance):
        return instance.agent.is_deleted

    get_is_deleted.short_description = _("is delete")


class AgentsRequestAttachmentInline(admin.TabularInline):
    verbose_name_plural = _("Attachments of the requests")
    fields = (
        "id",
        "get_file",
        "owner",
        "crm_uuid",
        "created_at",
        "is_deleted",
    )
    readonly_fields = (
        "id",
        "get_file",
        "owner",
        "crm_uuid",
        "created_at",
        "is_deleted",
    )
    model = AgentsRequestAttachment
    extra = 0

    def get_file(self, instance):
        """Returns the filename of the attached file."""
        if instance.attach_file:
            return mark_safe(
                '<a href="{}">{}</a>'.format(instance.attach_file.url, instance.attach_file.name)
            )
        else:
            return "No file attached"

    get_file.short_description = "file"


class ContractOfferInline(admin.TabularInline):
    verbose_name_plural = _("Contract offers")
    fields = (
        "id",
        "user",
        "number",
        "created_at",
        "is_deleted",
    )
    readonly_fields = (
        "id",
        "user",
        "number",
        "created_at",
        "is_deleted",
    )
    model = ContractOffer
    extra = 0


# ✅ ИСПОЛЬЗУЕТСЯ ОБЪЕДИНЕННЫЙ МОДУЛЬ django_admin_buttons.py 
# Все проблемы с кнопками и инлайнами решены в одном месте!

# ✅ ИСПОЛЬЗУЕМ ОБЪЕДИНЕННЫЙ МОДУЛЬ - ВСЕ ПРОБЛЕМЫ РЕШЕНЫ!
@admin.register(User)
class CustomUserAdmin(ExtraButtonsMixin, UserAdmin):
    inlines = (
        AgentsInline,
        AgentsRequestAttachmentInline,
    )
    list_display = (
        "email",
        "fio",
        "get_fio",
        "phone",
        "role",
        "is_active",
        "is_verified",
        "is_rejected",
        "is_need_new_password",
        "otp",
        "otp_expiry",
        "max_otp_try",
        "otp_max_out",
        "is_one_time_jwt_created",
        "date_registration",
        "exchange_type",
    )
    search_fields = (
        "email",
        "fio",
    )
    list_per_page = 25
    fieldsets = (
        (None, {"fields": ("email", "password")}),
        (_("Personal info"), {"fields": ("fio", "phone", "role")}),
        (
            _("Permissions"),
            {
                "fields": (
                    "is_active",
                    "is_staff",
                    "is_superuser",
                    "is_verified",
                    "is_rejected",
                    "is_need_new_password",
                    "exchange_type",
                    "groups",
                    "user_permissions",
                ),
            },
        ),
        (
            _("OTP"),
            {
                "fields": (
                    "otp",
                    "otp_expiry",
                    "max_otp_try",
                    "otp_max_out",
                )
            },
        ),
        (_("Other"), {"fields": ("is_one_time_jwt_created",)}),
        (
            _("Important dates"),
            {
                "fields": (
                    "last_login",
                    "date_joined",
                )
            },
        ),
    )
    readonly_fields = (
        "date_registration",
        "last_login",
        "date_joined",
    )
    add_form = UserCreationForm
    add_fieldsets = (
        (
            None,
            {
                "classes": ("wide",),
                "fields": ("email", "phone", "password1", "password2"),
            },
        ),
    )
    list_display_links = (
        "email",
        "fio",
    )
    list_filter = (
        "role",
        "is_active",
        "is_verified",
        "is_rejected",
        "is_one_time_jwt_created",
    )
    ordering = ("email",)
    actions = (
        "approve_registration",
        "reject_registration",
        "generate_otp",
    )

    def get_fio(self, obj):
        """Return user's full name for list display."""
        return obj.fio or f"{obj.first_name} {obj.last_name}".strip()
    
    get_fio.short_description = _("Full Name")

    @button(
        html_attrs={
            "style": "background-color:#28a745;color:white",
            "title": _("Employees report"),
        },
        change_list=True,
    )
    def generate_report(self, request):
        """Redirect to the employees report generation view."""
        url = reverse_lazy("admin:users_report")
        return redirect(url)

    def get_urls(self):
        from django.urls import path

        urls = super().get_urls()
        custom_urls = [
            path(
                "report/",
                self.admin_site.admin_view(self.users_report_view),
                name="users_report",
            ),
            path(
                "report/export/",
                self.admin_site.admin_view(self.export_users_report_xlsx),
                name="users_report_export",
            ),
        ]
        return custom_urls + urls

    @admin.action(description=_("Approve registration for selected users"))
    def approve_registration(self, request, queryset):
        for user in queryset:
            if user.is_active:
                self.message_user(request, _(f"User (id={user.id}) already activated"))
                continue
            try:
                UserApproveHelper.approve_user(user=user)
            except Exception as e:
                self.message_user(
                    request, _(f"Activation failed for user: {user.id} - {e}"), messages.ERROR
                )

    @admin.action(description=_("Reject registration for selected users"))
    def reject_registration(self, request, queryset):
        for user in queryset:
            user.is_active = False
            user.is_rejected = True
            user.save()
            email_body = (
                f"Добрый день, {user.fio}!\n"
                f'Администрацией отклонена регистрация на сайте "ЛК ПапаФинанс".\n'
            )

            Sender.send_mails_to_new_agent_verify(users_email=user.email, message=email_body)

    @admin.action(description=_("Generate new otp for selected users"))
    def generate_otp(self, request, queryset):
        for user in queryset:
            UserHelper(user).generate_otp()
            user.otp_max_out = None
            user.save()

    def _get_report_filters(self, request) -> dict:
        """Extract and validate report filters from request."""
        date_from = request.GET.get("date_from", "")
        date_to = request.GET.get("date_to", "")
        role = request.GET.get("role", "")
        agent_status = request.GET.get("agent_status", "")
        is_active = request.GET.get("is_active", "")
        is_verified = request.GET.get("is_verified", "")
        agent_inn = request.GET.get("agent_inn", "")  # Added agent INN filter

        try:
            page = int(request.GET.get("page", 1))
            if page < 1:
                page = 1
        except ValueError:
            page = 1

        return {
            "date_from": date_from,
            "date_to": date_to,
            "role": role,
            "agent_status": agent_status,
            "is_active": is_active,
            "is_verified": is_verified,
            "agent_inn": agent_inn,  # Added agent INN filter
            "page": page,
        }

    def _get_filtered_queryset(self, request, filters: dict):
        """
        Get base filtered queryset without expensive operations.

        Performance optimizations:
        - Filters applied early to reduce dataset size
        - Agent filters handled separately to avoid recursion
        - distinct() used only when necessary
        """
        qs = super().get_queryset(request)

        # Apply simple filters first
        if filters["date_from"]:
            try:
                qs = qs.filter(date_joined__date__gte=filters["date_from"])
            except (
                ValueError,
                ValidationError,
            ):  # FIXME: #002 - Invalid date handling in users admin report filters
                pass

        if filters["date_to"]:
            try:
                qs = qs.filter(date_joined__date__lte=filters["date_to"])
            except (
                ValueError,
                ValidationError,
            ):  # FIXME: #002 - Invalid date handling in users admin report filters
                pass

        if filters["role"]:
            qs = qs.filter(role=filters["role"])

        if filters["is_active"]:
            if filters["is_active"].lower() == "true":
                qs = qs.filter(is_active=True)
            elif filters["is_active"].lower() == "false":
                qs = qs.filter(is_active=False)

        if filters["is_verified"]:
            if filters["is_verified"].lower() == "true":
                qs = qs.filter(is_verified=True)
            elif filters["is_verified"].lower() == "false":
                qs = qs.filter(is_verified=False)

        # Handle agent-related filters separately to avoid recursion
        agent_user_ids = None

        if filters["agent_status"] or filters["agent_inn"]:
            # Build agent filter query separately
            agent_qs = Agent.objects.filter(is_deleted=False)

            if filters["agent_status"]:
                agent_qs = agent_qs.filter(status=filters["agent_status"])

            if filters["agent_inn"]:
                agent_qs = agent_qs.filter(inn__icontains=filters["agent_inn"])

            # Get user IDs through the agent queryset
            agent_user_ids = list(agent_qs.values_list("employees__id", flat=True).distinct())

            # Filter out None values
            agent_user_ids = [uid for uid in agent_user_ids if uid is not None]

            if agent_user_ids:
                qs = qs.filter(id__in=agent_user_ids)
            else:
                # No matching agents found, return empty queryset
                qs = qs.none()

        return qs

    def _get_report_data(self, request, filters: dict) -> tuple:
        """Get filtered and paginated report data."""
        # Get base filtered queryset
        qs = self._get_filtered_queryset(request, filters)

        # Apply ordering for consistent pagination
        qs = qs.order_by("email")

        # Simple pagination first
        paginator = Paginator(qs, per_page=20)

        try:
            page_obj = paginator.page(filters["page"])
        except (PageNotAnInteger, EmptyPage):
            page_obj = paginator.page(1)

        # Get user IDs from current page
        user_ids = [user.id for user in page_obj.object_list]

        if not user_ids:
            # Empty page, return as is
            return qs, paginator, page_obj

        # Create optimized queryset for current page only
        optimized_users = (
            self.model.objects.filter(id__in=user_ids)
            .select_related()
            .prefetch_related("agents")  # Simple prefetch without complex filters
            .order_by("email")
        )

        # Convert to list and add agent data manually
        users_list = []
        for user in optimized_users:
            # Get active agents for this user
            active_agents = [agent for agent in user.agents.all() if not agent.is_deleted]
            user.active_agents = active_agents
            user.agents_count = len(active_agents)
            users_list.append(user)

        # Replace page object list with optimized data
        page_obj.object_list = users_list

        return qs, paginator, page_obj

    def users_report_view(self, request):
        """Generate and display the users report."""
        if request.method != "GET":
            return HttpResponseNotAllowed(["GET"])

        filters = self._get_report_filters(request)
        all_report_data, paginator, page_obj = self._get_report_data(request, filters)
        report_generation_time = timezone.now()

        # Get unique agent INNs for dropdown filter
        agent_inns = (
            Agent.objects.filter(is_deleted=False)
            .values_list("inn", flat=True)
            .distinct()
            .order_by("inn")
        )

        context = {
            "title": _("Employees Report"),
            "report_data": page_obj.object_list if hasattr(page_obj, "object_list") else page_obj,
            "opts": self.model._meta,
            "filters": filters,
            "roles": UserRole.CHOICES,
            "agent_statuses": AgentsStatus.CHOICES,
            "agent_inns": agent_inns,
            "paginator": paginator,
            "page_obj": page_obj,
            "report_generation_time": report_generation_time,
        }

        return render(
            request,
            "admin/users/users_report.html",
            context,
        )

    def export_users_report_xlsx(self, request) -> HttpResponse:
        """Generate and export users report as XLSX file."""
        if request.method != "GET":
            return HttpResponseNotAllowed(["GET"])

        filters = self._get_report_filters(request)

        # Get filtered queryset using the same approach as report view
        qs = self._get_filtered_queryset(request, filters)

        # Apply ordering for consistent results
        qs = qs.order_by("email")

        # Use simple prefetch for all users in export
        optimized_qs = qs.select_related().prefetch_related("agents")

        # Create Excel workbook
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Employees Report"

        # Add headers
        headers = [
            gettext("ID"),
            gettext("Email"),
            gettext("FIO"),
            gettext("Phone"),
            gettext("Role"),
            gettext("Is Active"),
            gettext("Is Verified"),
            gettext("Is Rejected"),
            gettext("Agent INNs"),
            gettext("Agents Count"),
            gettext("Date Joined"),
        ]

        for col_num, header in enumerate(headers, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.value = str(header)
            cell.font = openpyxl.styles.Font(bold=True)

        # Add data with proper Unicode handling
        row_num = 2
        for user in optimized_qs.iterator(chunk_size=100):  # Use iterator for memory efficiency
            try:
                worksheet.cell(row=row_num, column=1).value = user.id
                worksheet.cell(row=row_num, column=2).value = str(user.email) if user.email else ""
                worksheet.cell(row=row_num, column=3).value = str(user.fio) if user.fio else ""
                worksheet.cell(row=row_num, column=4).value = str(user.phone) if user.phone else ""

                role_value = dict(UserRole.CHOICES).get(user.role, user.role)
                worksheet.cell(row=row_num, column=5).value = str(role_value)

                worksheet.cell(row=row_num, column=6).value = (
                    _("Yes") if user.is_active else _("No")
                )
                worksheet.cell(row=row_num, column=7).value = (
                    _("Yes") if user.is_verified else _("No")
                )
                worksheet.cell(row=row_num, column=8).value = (
                    _("Yes") if user.is_rejected else _("No")
                )

                # Get agent INNs and count manually without complex filtering
                active_agents = [agent for agent in user.agents.all() if not agent.is_deleted]
                agent_inns = [str(agent.inn) for agent in active_agents if agent.inn]
                worksheet.cell(row=row_num, column=9).value = (
                    ", ".join(agent_inns) if agent_inns else ""
                )
                worksheet.cell(row=row_num, column=10).value = len(active_agents)
                worksheet.cell(row=row_num, column=11).value = user.date_joined.strftime(
                    "%Y-%m-%d %H:%M:%S"
                )

                row_num += 1

            except Exception as e:
                logger.error(f"Error processing user {user.id} in Excel export: {e}")
                continue

        # Set column widths
        for col_num in range(1, len(headers) + 1):
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 15

        # Create response with Excel file
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
        filename = f"employees_report_{timestamp}.xlsx"
        response["Content-Disposition"] = f"attachment; filename={filename}"

        # Save workbook to response
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        response.write(buffer.getvalue())
        buffer.close()

        return response


@admin.register(BlackListedAccessToken)
class BlackListedAccessTokenAdmin(admin.ModelAdmin):
    list_display = (
        "id",
        "jti",
        "jti_refresh",
        "user",
        "created_at",
        "expires_at",
        "blacklisted_at",
        "timestamp",
    )
    list_display_links = (
        "id",
        "jti",
    )
    search_fields = (
        "user__id",
        "user__email",
        "jti",
    )
    ordering = (
        "user",
        "-created_at",
    )

 